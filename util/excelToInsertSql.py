#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import datetime
import decimal
import os
from typing import Any, List

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert Excel file to MySQL INSERT SQL."
    )
    parser.add_argument(
        "-i", "--input",
        required=True,
        help="Input Excel file path (e.g. test.xlsx)"
    )
    parser.add_argument(
        "-t", "--table",
        required=True,
        help="MySQL table name"
    )
    parser.add_argument(
        "-o", "--output",
        default=None,
        help="Output SQL file path (default: <input>.sql)"
    )
    parser.add_argument(
        "-s", "--sheet",
        default=None,
        help="Sheet name to read (default: first sheet)"
    )
    parser.add_argument(
        "-b", "--batch-size",
        type=int,
        default=500,
        help="Number of rows per INSERT statement (default: 500)"
    )
    return parser.parse_args()


def sql_escape_value(value: Any) -> str:
    """Convert Python value to SQL literal."""
    if value is None:
        return "NULL"

    # Empty string or only whitespace当作空字符串，不是 NULL
    if isinstance(value, str):
        # 先保留用户的空字符串语义
        v = value
        # 转义单引号
        v = v.replace("'", "''")
        return f"'{v}'"

    if isinstance(value, (int, float, decimal.Decimal)):
        return str(value)

    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return f"'{value.isoformat(sep=' ', timespec='seconds')}'"

    if isinstance(value, bool):
        return "1" if value else "0"

    # 其他类型统一转成字符串
    v = str(value).replace("'", "''")
    return f"'{v}'"


def build_insert_statement(
    table_name: str,
    columns: List[str],
    rows: List[List[Any]]
) -> str:
    """Build one INSERT statement with multiple VALUES."""
    col_list = ", ".join(f"`{c}`" for c in columns)
    values_parts = []

    for row in rows:
        # 保证行长度与列数对齐：多余截断，不足补 None
        row = (row + [None] * len(columns))[: len(columns)]
        literals = [sql_escape_value(v) for v in row]
        values_parts.append("(" + ", ".join(literals) + ")")

    values_str = ",\n  ".join(values_parts)
    sql = f"INSERT INTO `{table_name}` ({col_list}) VALUES\n  {values_str};\n"
    return sql


def main() -> None:
    args = parse_args()

    input_path = args.input
    if not os.path.isfile(input_path):
        raise FileNotFoundError(f"Input file not found: {input_path}")

    output_path = args.output
    if output_path is None:
        base, _ = os.path.splitext(input_path)
        output_path = base + ".sql"

    table_name = args.table
    batch_size = max(1, args.batch_size)

    # 只读模式，数据量大时稍微省内存一点
    wb = load_workbook(filename=input_path, read_only=True, data_only=True)

    if args.sheet:
        if args.sheet not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{args.sheet}' not found in workbook. "
                f"Available sheets: {wb.sheetnames}"
            )
        ws = wb[args.sheet]
    else:
        ws = wb[wb.sheetnames[0]]

    # 打开输出文件
    with open(output_path, "w", encoding="utf-8") as f_out:
        # 迭代行：min_row=1 表示从第一行开始
        rows_iter = ws.iter_rows(values_only=True)

        # 第一行是字段名
        try:
            header = next(rows_iter)
        except StopIteration:
            print("Excel file is empty, nothing to convert.")
            return

        if header is None:
            print("Header row is empty, abort.")
            return

        columns = [str(c).strip() if c is not None else "" for c in header]
        # 去掉空字段名的列
        valid_indexes = [i for i, c in enumerate(columns) if c]
        columns = [columns[i] for i in valid_indexes]

        if not columns:
            print("No valid column names found in header, abort.")
            return

        batch_rows: List[List[Any]] = []
        total_rows = 0
        for row in rows_iter:
            if row is None:
                continue
            # 只取有有效字段名的列
            row_values = [row[i] if i < len(row) else None for i in valid_indexes]

            # 判断该行是否全空，如果全空就跳过
            if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row_values):
                continue

            batch_rows.append(row_values)
            total_rows += 1

            if len(batch_rows) >= batch_size:
                sql = build_insert_statement(table_name, columns, batch_rows)
                f_out.write(sql)
                batch_rows.clear()

        # 写出最后一个不满 batch 的 batch
        if batch_rows:
            sql = build_insert_statement(table_name, columns, batch_rows)
            f_out.write(sql)

    print(f"Done. Total rows converted: {total_rows}")
    print(f"SQL written to: {output_path}")


if __name__ == "__main__":
    main()